"""Gera o artigo científico no padrão Unisales (docs/artigo/Artigo_DW_AdventureWorks.docx).

Os números citados no texto são lidos do Data Warehouse no momento da geração.
Uso: python scripts/gerar_artigo.py --github https://github.com/<usuario>/adventureworks-dw
"""
import argparse
import re
import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageOps

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "scripts"))
from etl.config import dw_engine  # noqa: E402
from gerar_dicionario import TABELAS, carregar_dicionario  # noqa: E402

FIG = RAIZ / "docs" / "figuras"
SAIDA = RAIZ / "docs" / "artigo" / "Artigo_DW_AdventureWorks.docx"
FONTE = "Arial"
DATA_ACESSO = "23 set. 2026"


# ----------------------------------------------------------------------------- formatação
def br(v, casas=2):
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def moeda(v):
    return f"US$ {br(v / 1e6, 1)} milhões" if abs(v) >= 1e6 else f"US$ {br(v)}"


def fmt_par(p, tamanho=12, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, depois=6,
            entrelinhas=1.0, recuo_esq=None):
    pf = p.paragraph_format
    pf.alignment = alinhamento
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.line_spacing = entrelinhas
    pf.first_line_indent = Cm(0)
    if recuo_esq is not None:
        pf.left_indent = Cm(recuo_esq)
    for r in p.runs:
        r.font.name = FONTE
        r.font.size = Pt(tamanho)
    return p


def add_runs(p, texto, tamanho=12):
    """Aceita **negrito** e *itálico* inline."""
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = p.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("*"):
            r = p.add_run(parte[1:-1]); r.italic = True
        else:
            r = p.add_run(parte)
        r.font.name = FONTE
        r.font.size = Pt(tamanho)
    return p


class Artigo:
    def __init__(self):
        self.doc = Document()
        self.n_fig = self.n_quadro = self.n_tabela = 0
        self._configurar()

    def _configurar(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONTE
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.left_margin = Cm(3), Cm(3)
        sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
        # Paginação: Arial 10, canto superior direito
        p = sec.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONTE); fonts.set(qn("w:hAnsi"), FONTE)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "20")
        rpr.append(fonts); rpr.append(sz); r.append(rpr)
        t = OxmlElement("w:t"); t.text = "1"; r.append(t)
        fld.append(r)
        p._p.append(fld)

    # --- blocos de texto
    def par(self, texto, **kw):
        tamanho = kw.pop("tamanho", 12)
        p = self.doc.add_paragraph()
        add_runs(p, texto, tamanho)
        return fmt_par(p, tamanho=tamanho, **kw)

    def vazio(self, entrelinhas=1.0):
        p = self.doc.add_paragraph()
        return fmt_par(p, antes=0, depois=0, entrelinhas=entrelinhas)

    def titulo(self, texto, nivel=1):
        p = self.doc.add_paragraph()
        r = p.add_run(texto.upper() if nivel <= 2 else texto)
        r.bold = nivel in (1, 3)
        fmt_par(p, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, antes=12 if nivel == 1 else 6, depois=6)
        p.paragraph_format.keep_with_next = True
        return p

    def lista(self, itens, tamanho=12):
        for i, item in enumerate(itens):
            letra = chr(ord("a") + i)
            fim = ";" if i < len(itens) - 1 else "."
            p = self.doc.add_paragraph()
            add_runs(p, f"{letra}) {item}{fim}", tamanho)
            fmt_par(p, tamanho=tamanho, depois=3, recuo_esq=0.6)

    # --- legendas
    def _legenda(self, rotulo, titulo):
        p = self.doc.add_paragraph()
        add_runs(p, f"{rotulo} – {titulo}", 12)
        fmt_par(p, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, depois=0)
        p.paragraph_format.keep_with_next = True

    def _fonte(self, texto="Elaboração própria (2026)."):
        p = self.doc.add_paragraph()
        add_runs(p, f"Fonte: {texto}", 10)
        fmt_par(p, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=0)
        self.vazio(entrelinhas=1.5)

    def figura(self, arquivo, titulo, largura_cm=15.5, fonte=None):
        self.n_fig += 1
        self.vazio(entrelinhas=1.5)
        self._legenda(f"Figura {self.n_fig}", titulo)
        img = ImageOps.expand(Image.open(FIG / arquivo).convert("RGB"), border=3, fill="black")
        buf = BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
        p = self.doc.add_paragraph()
        p.add_run().add_picture(buf, width=Cm(largura_cm))
        fmt_par(p, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, depois=0)
        p.paragraph_format.keep_with_next = True
        self._fonte(fonte or "Elaboração própria (2026), a partir do Data Warehouse.")
        return self.n_fig

    # --- tabelas e quadros
    def _grade(self, cabecalho, linhas, larguras, fechado, fonte_mono=False):
        t = self.doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borda = "single"
        bordas = {"top": borda, "bottom": borda, "insideH": borda if fechado else "nil",
                  "left": borda if fechado else "nil", "right": borda if fechado else "nil",
                  "insideV": borda if fechado else "nil"}
        tblpr = t._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        for lado, val in bordas.items():
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), val); b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "000000")
            el.append(b)
        tblpr.append(el)
        for i, linha in enumerate([cabecalho] + linhas):
            for j, valor in enumerate(linha):
                cel = t.cell(i, j)
                cel.width = Cm(larguras[j])
                p = cel.paragraphs[0]
                r = p.add_run(str(valor))
                r.font.size = Pt(10)
                r.font.name = "Courier New" if (fonte_mono and i > 0) else FONTE
                r.bold = i == 0
                fmt_par(p, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=0)
                for r in p.runs:
                    if fonte_mono and i > 0:
                        r.font.name = "Courier New"
                if i == 0 and not fechado:  # traço separando cabeçalho (tabela aberta)
                    tcpr = cel._tc.get_or_add_tcPr()
                    tb = OxmlElement("w:tcBorders")
                    b = OxmlElement("w:bottom")
                    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "000000")
                    tb.append(b); tcpr.append(tb)
        # Repete o cabeçalho quando a tabela passa de uma página
        trpr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trpr.append(h)
        return t

    def tabela(self, titulo, cabecalho, linhas, larguras, fonte=None):
        self.n_tabela += 1
        self.vazio(entrelinhas=1.5)
        self._legenda(f"Tabela {self.n_tabela}", titulo)
        self._grade(cabecalho, linhas, larguras, fechado=False)
        self._fonte(fonte or "Elaboração própria (2026), a partir do Data Warehouse.")
        return self.n_tabela

    def quadro(self, titulo, cabecalho, linhas, larguras, fonte=None, mono=False):
        self.n_quadro += 1
        self.vazio(entrelinhas=1.5)
        self._legenda(f"Quadro {self.n_quadro}", titulo)
        self._grade(cabecalho, linhas, larguras, fechado=True, fonte_mono=mono)
        self._fonte(fonte or "Elaboração própria (2026).")
        return self.n_quadro

    def quadro_sql(self, titulo, sql):
        """Quadro fechado de uma célula contendo o código SQL em fonte monoespaçada."""
        self.n_quadro += 1
        self.vazio(entrelinhas=1.5)
        self._legenda(f"Quadro {self.n_quadro}", titulo)
        t = self._grade(["SQL"], [[""]], [15.5], fechado=True)
        cel = t.cell(1, 0)
        cel.paragraphs[0]._p.getparent().remove(cel.paragraphs[0]._p)
        for linha in sql.strip("\n").splitlines():
            p = cel.add_paragraph()
            r = p.add_run(linha if linha else " ")
            r.font.name = "Courier New"; r.font.size = Pt(8)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        self._fonte()
        return self.n_quadro

    def salvar(self):
        SAIDA.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(SAIDA)
        print("  artigo:", SAIDA.relative_to(RAIZ))


# ----------------------------------------------------------------------------- dados
def kpi_sql():
    """Separa o arquivo de views em {nome_view: sql}."""
    texto = (RAIZ / "sql" / "02_kpis_views.sql").read_text(encoding="utf-8")
    blocos = {}
    for m in re.finditer(r"(-- KPI \d+ - [^\n]+\n)(CREATE OR REPLACE VIEW (\w+) AS\n.*?;)", texto, re.S):
        blocos[m.group(3)] = m.group(2)
    return blocos


def carregar_numeros(eng):
    q = lambda sql: pd.read_sql(sql, eng)  # noqa: E731
    n = {}
    n["resumo"] = q("SELECT * FROM dw.vw_kpi_resumo").iloc[0]
    anual = q("SELECT ano, SUM(receita_liquida) r FROM dw.vw_kpi01_receita_liquida GROUP BY ano ORDER BY ano")
    n["anual"] = anual.set_index("ano")["r"]
    canal = q("SELECT ano, canal_venda, receita_liquida FROM dw.vw_kpi03_ticket_medio")
    n["canal"] = canal.pivot(index="ano", columns="canal_venda", values="receita_liquida")
    n["ticket"] = q("SELECT * FROM dw.vw_kpi03_ticket_medio ORDER BY ano, canal_venda")
    n["margem"] = q("SELECT * FROM dw.vw_kpi04_margem_bruta ORDER BY ano, canal_venda")
    n["margem_canal"] = q("""SELECT canal_venda, ROUND(100*SUM(lucro_bruto)/SUM(valor_liquido),2) m,
                                    SUM(valor_liquido) r FROM dw.fato_vendas GROUP BY 1""").set_index("canal_venda")
    n["yoy"] = q("SELECT * FROM dw.vw_kpi05_crescimento_yoy ORDER BY ano, mes")
    n["terr"] = q("SELECT * FROM dw.vw_kpi06_receita_territorio ORDER BY receita_liquida DESC")
    n["grupo"] = q("""SELECT grupo, SUM(participacao_pct) p FROM dw.vw_kpi06_receita_territorio
                      GROUP BY grupo ORDER BY p DESC""")
    n["cat"] = q("SELECT * FROM dw.vw_kpi07_receita_categoria ORDER BY receita_liquida DESC")
    n["desc"] = q("SELECT * FROM dw.vw_kpi08_desconto ORDER BY ano, canal_venda")
    n["recompra"] = q("SELECT * FROM dw.vw_kpi09_taxa_recompra").set_index("tipo_cliente")
    n["cota"] = q("""SELECT ano, ROUND(AVG(atingimento_pct),1) media,
                            COUNT(*) FILTER (WHERE atingimento_pct >= 100) bateram, COUNT(*) total
                     FROM dw.vw_kpi10_atingimento_cota GROUP BY ano ORDER BY ano""")
    n["linhas"] = q("SELECT tabela, linhas FROM dw.etl_execucao ORDER BY id DESC LIMIT 9").set_index("tabela")["linhas"]
    n["periodo"] = q("SELECT MIN(t.data) i, MAX(t.data) f FROM dw.fato_vendas f JOIN dw.dim_tempo t ON t.sk_tempo=f.sk_data_pedido").iloc[0]
    n["fim_revenda"] = q("""SELECT MAX(t.data) d FROM dw.fato_vendas f JOIN dw.dim_tempo t ON t.sk_tempo=f.sk_data_pedido
                            WHERE canal_venda='Revenda'""").iloc[0]["d"]
    return n


# ----------------------------------------------------------------------------- conteúdo
def escrever(a: Artigo, n, github):
    r = n["resumo"]
    anual, canal = n["anual"], n["canal"]
    cresc_24 = 100 * (anual[2024] - anual[2023]) / anual[2023]
    share_online = lambda ano: 100 * canal.loc[ano, "Online"] / canal.loc[ano].sum()  # noqa: E731
    mc = n["margem_canal"]
    cat = n["cat"].set_index("categoria")
    terr = n["terr"]
    top3 = terr.head(3)["participacao_pct"].sum()
    grupos = n["grupo"].set_index("grupo")["p"]
    rec = n["recompra"]
    cota = n["cota"].set_index("ano")
    tk = n["ticket"]
    tk_on = tk[tk.canal_venda == "Online"].set_index("ano")["ticket_medio"]
    tk_rv = tk[tk.canal_venda == "Revenda"].set_index("ano")["ticket_medio"]
    m = n["margem"]
    m_rv = m[m.canal_venda == "Revenda"].set_index("ano")["margem_bruta_pct"]
    periodo_i, periodo_f = n["periodo"]["i"], n["periodo"]["f"]

    # ---------------- Pré-textuais
    t = a.doc.add_paragraph()
    add_runs(t, "**DATA WAREHOUSE PARA ANÁLISE DE VENDAS DA ADVENTUREWORKS: MODELAGEM MULTIDIMENSIONAL, ETL EM PYTHON E INDICADORES DE DESEMPENHO**")
    fmt_par(t, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, depois=12)
    for linha in ["Nome Completo do(a) Aluno(a)¹", "Nome Completo do(a) Orientador(a)²"]:
        a.par(linha, alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, depois=0)
    a.vazio()

    a.par("**RESUMO**", alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    a.par(
        "Organizações acumulam grandes volumes de dados transacionais, mas a estrutura normalizada "
        "desses sistemas dificulta a análise gerencial. Este artigo apresenta a construção de um Data "
        "Warehouse para a área de vendas da empresa fictícia AdventureWorks, disponibilizada pela "
        "Microsoft como base de exemplo. Avaliou-se o modelo transacional de origem, com 68 tabelas "
        "distribuídas em cinco esquemas, e propôs-se um modelo multidimensional em esquema estrela "
        "composto por duas tabelas fato e sete dimensões. Uma rotina de extração, transformação e "
        f"carga (ETL) escrita em Python popula o Data Warehouse em PostgreSQL, carregando {br(n['linhas']['fato_vendas'], 0)} "
        "itens de pedido com custo histórico, rateio de frete e impostos e chaves substitutas. Foram "
        "definidos dez indicadores de desempenho, implementados como visões SQL e consumidos por um "
        f"dashboard. Os resultados mostram receita líquida de {moeda(r.receita_liquida)} no período, "
        f"crescimento de {br(cresc_24, 1)}% entre 2023 e 2024 e forte contraste de rentabilidade entre "
        f"canais: margem bruta de {br(mc.loc['Online','m'], 1)}% nas vendas online contra "
        f"{br(mc.loc['Revenda','m'], 1)}% na revenda. Conclui-se que a modelagem dimensional, aliada a "
        "uma ETL reprodutível, transforma dados operacionais em informação acionável para a tomada de decisão.",
        tamanho=12)
    a.par("**Palavras-chave:** data warehouse; modelagem dimensional; ETL; indicadores de desempenho; PostgreSQL.")
    a.vazio()
    a.par("**ABSTRACT**", alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    a.par(
        "Organizations accumulate large volumes of transactional data, but the normalized structure of "
        "these systems makes managerial analysis difficult. This paper presents the construction of a "
        "Data Warehouse for the sales area of the fictitious company AdventureWorks, a sample database "
        "provided by Microsoft. The source transactional model, with 68 tables in five schemas, was "
        "evaluated, and a star schema with two fact tables and seven dimensions was proposed. An "
        "extract, transform and load (ETL) routine written in Python populates the Data Warehouse in "
        f"PostgreSQL, loading {br(n['linhas']['fato_vendas'], 0).replace('.', ',')} order lines with historical cost, "
        "allocated freight and taxes and surrogate keys. Ten key performance indicators were defined, "
        "implemented as SQL views and consumed by a dashboard. Results show net revenue of "
        f"US$ {r.receita_liquida / 1e6:,.1f} million, {cresc_24:.1f}% growth between 2023 and 2024 and a "
        f"strong profitability gap between channels: {mc.loc['Online','m']:.1f}% gross margin online versus "
        f"{mc.loc['Revenda','m']:.1f}% in the reseller channel. Dimensional modeling combined with a "
        "reproducible ETL turns operational data into actionable information for decision making.")
    a.par("**Keywords:** data warehouse; dimensional modeling; ETL; key performance indicators; PostgreSQL.")

    # ---------------- 1 Introdução
    a.titulo("1 INTRODUÇÃO")
    a.par(
        "Os sistemas transacionais, também chamados de sistemas OLTP (Online Transaction Processing), são "
        "projetados para registrar operações do dia a dia com rapidez e integridade. Para isso, seus "
        "bancos de dados são altamente normalizados, o que evita redundâncias, mas espalha a informação "
        "por dezenas de tabelas. Quando um gestor precisa saber qual região é mais lucrativa ou se a "
        "equipe comercial está cumprindo suas metas, essa estrutura exige consultas extensas, lentas e "
        "sujeitas a erro.")
    a.par(
        "O Data Warehouse (DW) surge como resposta a esse problema. Inmon (2005) o define como uma coleção "
        "de dados orientada por assunto, integrada, não volátil e variável no tempo, destinada a apoiar a "
        "tomada de decisão. Kimball e Ross (2013) propõem a modelagem dimensional, cujo principal artefato "
        "é o esquema estrela: tabelas fato, que armazenam as medidas de um processo de negócio, cercadas "
        "por tabelas dimensão, que fornecem o contexto (quem, o quê, onde e quando) para analisá-las.")
    a.par(
        "Entre a origem transacional e o DW está o processo de extração, transformação e carga (ETL), "
        "responsável por ler os dados, limpá-los, integrá-los e gravá-los no formato dimensional. Segundo "
        "Kimball e Caserta (2004), a ETL costuma consumir a maior parte do esforço de um projeto de DW, "
        "pois é nela que se garantem a qualidade e a consistência da informação entregue ao usuário.")
    a.par(
        "Este trabalho tem como objetivo geral construir um Data Warehouse para a área de vendas da "
        "AdventureWorks, empresa fictícia fabricante de bicicletas cuja base de dados é distribuída pela "
        "Microsoft (2024) para fins didáticos. Os objetivos específicos são:")
    a.lista([
        "avaliar o modelo de dados transacional disponível",
        "elaborar dez indicadores de desempenho (KPIs) para a gestão comercial",
        "propor um modelo multidimensional e seu dicionário de dados",
        "desenvolver uma ETL em Python que popule o DW em PostgreSQL",
        "implementar os indicadores em SQL e apresentá-los em um dashboard, aplicando técnicas de storytelling com dados",
    ])
    a.par(
        "O artigo está organizado da seguinte forma: a seção 2 apresenta o desenvolvimento, da avaliação da "
        "origem à análise dos indicadores, e a seção 3 traz as considerações finais.")

    # ---------------- 2 Desenvolvimento
    a.titulo("2 DESENVOLVIMENTO")
    a.par(
        "O projeto foi executado em um ambiente local com PostgreSQL 16, que hospeda tanto a base de origem "
        "quanto o DW, e Python 3.9 com as bibliotecas pandas, SQLAlchemy e psycopg2. A base AdventureWorks, "
        "originalmente em SQL Server, foi instalada no PostgreSQL por meio dos scripts de conversão "
        "publicados por Thwaits (2024), que reaproveitam os arquivos CSV oficiais da Microsoft.")

    a.titulo("2.1 AVALIAÇÃO DO MODELO DE DADOS DE ORIGEM", 2)
    a.par(
        "O banco AdventureWorks possui 68 tabelas organizadas em cinco esquemas: *person* (pessoas, "
        "endereços e contatos), *humanresources* (funcionários e departamentos), *production* (produtos, "
        "categorias, custos, estoque e ordens de produção), *purchasing* (fornecedores, compras e métodos "
        "de envio) e *sales* (clientes, lojas, vendedores, territórios, pedidos e promoções). O modelo está "
        "na terceira forma normal e usa a tabela *person.businessentity* como supertipo de pessoas, lojas e "
        "fornecedores.")
    a.par(
        f"O processo de negócio escolhido foi a venda, por ser o mais relevante para a gestão e o de maior "
        f"volume: {br(n['linhas']['fato_vendas'], 0)} itens de pedido (*sales.salesorderdetail*) ligados a "
        f"{br(r.qtd_pedidos, 0)} cabeçalhos de pedido (*sales.salesorderheader*), entre "
        f"{periodo_i:%d/%m/%Y} e {periodo_f:%d/%m/%Y}. A avaliação identificou as seguintes características "
        "que orientaram o projeto:")
    a.lista([
        "a hierarquia de produto está normalizada em três tabelas (*product*, *productsubcategory* e "
        "*productcategory*), e cerca de 40% dos produtos não possuem subcategoria, pois são componentes internos",
        "o cliente pode ser uma pessoa física (venda online) ou uma loja (revenda), identificados pelas colunas "
        "*personid* e *storeid* de *sales.customer*, e seu nome e endereço estão em outros esquemas",
        "o item de pedido não armazena custo; o custo precisa ser obtido em *production.productcosthistory*, "
        "que guarda o custo padrão por período de vigência",
        "frete (*freight*) e impostos (*taxamt*) existem apenas no cabeçalho do pedido",
        "pedidos online não têm vendedor (*salespersonid* nulo), e as cotas dos vendedores estão em "
        "*sales.salespersonquotahistory*, com periodicidade trimestral",
        "diversos atributos usam códigos (por exemplo, *productline* = R, M, T, S), pouco legíveis para o usuário final",
    ])

    a.titulo("2.2 INDICADORES PROPOSTOS", 2)
    a.par(
        "Os indicadores foram escolhidos para responder a quatro perguntas de negócio: quanto vendemos, "
        "quanto ganhamos, onde e o quê vendemos, e como se comportam clientes e vendedores. O Quadro 1 "
        "apresenta os dez indicadores, sua fórmula de cálculo e a pergunta de negócio que cada um responde.")
    a.quadro("Indicadores de desempenho propostos",
             ["Nº", "Indicador", "Fórmula", "Pergunta de negócio"],
             [
                 ["1", "Receita líquida", "Σ valor_liquido", "Quanto a empresa faturou por mês e canal?"],
                 ["2", "Volume de pedidos", "COUNT(DISTINCT numero_pedido); Σ quantidade", "Quantos pedidos e unidades foram vendidos?"],
                 ["3", "Ticket médio", "Receita líquida ÷ nº de pedidos", "Quanto vale, em média, cada pedido?"],
                 ["4", "Margem bruta %", "(Receita − custo) ÷ receita × 100", "Qual canal é mais rentável?"],
                 ["5", "Crescimento YoY %", "(Receita mês − mesmo mês ano anterior) ÷ ano anterior × 100", "O negócio está crescendo?"],
                 ["6", "Receita por território", "Σ valor_liquido por território; participação %", "Onde estão os mercados mais relevantes?"],
                 ["7", "Receita e margem por categoria", "Σ valor_liquido e margem por categoria", "Quais linhas de produto sustentam o negócio?"],
                 ["8", "Desconto médio %", "Σ valor_desconto ÷ Σ valor_bruto × 100", "Quanto se abre mão de receita com promoções?"],
                 ["9", "Taxa de recompra %", "Clientes com > 1 pedido ÷ clientes ativos × 100", "Os clientes voltam a comprar?"],
                 ["10", "Atingimento de cota %", "Vendas do vendedor no trimestre ÷ cota × 100", "A equipe comercial cumpre as metas?"],
             ], [1.0, 3.2, 5.8, 5.5])

    a.titulo("2.3 MODELO MULTIDIMENSIONAL PROPOSTO", 2)
    a.par(
        "Seguindo o processo de quatro passos de Kimball e Ross (2013), foram definidos: (1) o processo de "
        "negócio, vendas; (2) o grão, um item de pedido de venda, o nível mais detalhado disponível; (3) as "
        "dimensões que descrevem cada item; e (4) os fatos numéricos medidos. Como a análise de cotas tem "
        "grão diferente (vendedor por trimestre), criou-se uma segunda tabela fato que compartilha as "
        "dimensões de vendedor e tempo com a primeira, formando uma constelação de fatos com dimensões "
        "conformadas. A Figura 1 apresenta o modelo.")
    a.figura("fig1_modelo_estrela.png", "Modelo multidimensional (esquema estrela) proposto", largura_cm=15.5)
    a.par("As principais decisões de modelagem foram:")
    a.lista([
        "**chaves substitutas** (sk_*) em todas as dimensões, isolando o DW das chaves da origem, que são mantidas como chave natural (NK)",
        "**dim_tempo** gerada pela ETL com uma linha por dia e chave no formato AAAAMMDD, usada em dois papéis na fato de vendas (data do pedido e data de envio)",
        "**dim_produto desnormalizada**, reunindo produto, modelo, subcategoria e categoria em uma única tabela, com códigos traduzidos para o português",
        "**dim_cliente unificada** para pessoas físicas e lojas, com o atributo tipo_cliente",
        "**membro especial** na dim_vendedor (vendedor_id = 0, \"Venda online\") para evitar chaves nulas na fato",
        "**dimensões degeneradas** numero_pedido e canal_venda mantidas na própria fato, pois não possuem outros atributos",
        "**fatos aditivos** (valores e quantidades) armazenados na fato; razões como margem e ticket médio são calculadas nas consultas, para não somar percentuais",
    ])

    a.titulo("2.4 DICIONÁRIO DE DADOS", 2)
    colunas, desc_tabelas = carregar_dicionario()
    a.par(
        "O dicionário de dados foi registrado no próprio catálogo do PostgreSQL por meio de comandos "
        "COMMENT ON, no script de criação do DW, e extraído automaticamente para os quadros a seguir. "
        "Na coluna Chave, PK indica chave primária e FK, chave estrangeira.")
    for tab in TABELAS:
        linhas = [[c.coluna, c.tipo.replace("character varying", "varchar").replace("timestamp without time zone", "timestamp"),
                   c.chave, c.descricao or ""] for c in colunas[colunas.tabela == tab].itertuples()]
        a.quadro(f"Dicionário de dados da tabela {tab} ({desc_tabelas.get(tab, '').lower()})",
                 ["Coluna", "Tipo", "Chave", "Descrição"], linhas, [3.8, 2.8, 1.4, 7.5])

    a.titulo("2.5 DESCRIÇÃO DA ETL", 2)
    a.par(
        "A ETL foi desenvolvida em Python, organizada em módulos com responsabilidades separadas e "
        "orquestrada pelo script run_etl.py. As conexões são configuradas por variáveis de ambiente "
        "(arquivo .env), o que permite executar o mesmo código contra outros servidores. A estratégia "
        "adotada é a carga completa (*full load*): a cada execução o schema dw é recriado e todas as "
        "tabelas são recarregadas, o que é adequado ao volume do projeto e garante reprodutibilidade. O "
        "Quadro 11 resume as etapas.")
    a.quadro("Etapas da ETL",
             ["Etapa", "Módulo", "O que faz"],
             [
                 ["Extração", "etl/extract.py",
                  "Executa dez consultas SQL na origem (produtos, custos, territórios, clientes, endereços, vendedores, "
                  "promoções, métodos de envio, cotas e vendas) e carrega o resultado em DataFrames do pandas."],
                 ["Transformação", "etl/transform.py",
                  "Gera a dim_tempo; desnormaliza a hierarquia de produto; traduz códigos (linha, classe, estilo); "
                  "substitui nulos por \"Não informado\"; unifica pessoas e lojas na dim_cliente; cria o membro "
                  "\"Venda online\"; atribui chaves substitutas; busca o custo vigente na data do pedido; calcula "
                  "valores bruto, desconto, líquido, custo e lucro; rateia frete e impostos pela participação do "
                  "item no subtotal; resolve as chaves estrangeiras e interrompe a carga se alguma não for encontrada."],
                 ["Carga", "etl/load.py",
                  "Recria o schema a partir de sql/01_create_dw.sql, carrega dimensões e depois fatos com o comando "
                  "COPY do PostgreSQL (carga em massa), ajusta as sequences e registra cada tabela em dw.etl_execucao."],
                 ["Publicação", "run_etl.py",
                  "Cria as views dos indicadores (sql/02_kpis_views.sql) e exibe a receita total carregada para conferência."],
             ], [2.6, 3.0, 9.9])
    a.par(
        "Duas regras de transformação merecem destaque. A primeira é o **custo histórico**: para cada item, "
        "a ETL localiza em *productcosthistory* o custo padrão vigente na data do pedido e, na ausência de "
        "histórico, utiliza o custo atual do produto. Isso evita que a margem de vendas antigas seja "
        "distorcida por reajustes posteriores. A segunda é o **rateio de frete e impostos**, que distribui os "
        "valores do cabeçalho entre os itens na proporção do valor líquido de cada um, permitindo analisá-los "
        "por produto ou categoria.")
    a.par(
        f"A execução completa leva cerca de 13 segundos. A validação consistiu em comparar a receita "
        f"carregada no DW ({moeda(r.receita_liquida)}) com a soma da coluna *subtotal* de "
        "*sales.salesorderheader* na origem, e os valores coincidiram. A Tabela 1 apresenta o volume carregado por tabela.")
    a.tabela("Linhas carregadas por tabela do Data Warehouse",
             ["Tabela", "Tipo", "Linhas"],
             [[t, "Fato" if t.startswith("fato") else "Dimensão", br(n["linhas"][t], 0)] for t in TABELAS],
             [5.5, 3.0, 3.0])
    p = a.doc.add_paragraph()
    add_runs(p, f"O código-fonte completo, com instruções de instalação e execução, está disponível em: {github}.")
    fmt_par(p)

    a.titulo("2.6 IMPLEMENTAÇÃO DOS INDICADORES EM SQL", 2)
    a.par(
        "Cada indicador foi implementado como uma view no schema dw, de modo que a regra de cálculo fique "
        "centralizada no banco e qualquer ferramenta de visualização, como Power BI ou Metabase, obtenha os "
        "mesmos números. Os quadros a seguir apresentam o SQL de cada indicador e a Tabela 2 mostra o "
        "resultado obtido.")
    nomes = {
        "vw_kpi01_receita_liquida": "Receita líquida", "vw_kpi02_volume_pedidos": "Volume de pedidos",
        "vw_kpi03_ticket_medio": "Ticket médio", "vw_kpi04_margem_bruta": "Margem bruta",
        "vw_kpi05_crescimento_yoy": "Crescimento YoY", "vw_kpi06_receita_territorio": "Receita por território",
        "vw_kpi07_receita_categoria": "Receita e margem por categoria", "vw_kpi08_desconto": "Desconto médio",
        "vw_kpi09_taxa_recompra": "Taxa de recompra", "vw_kpi10_atingimento_cota": "Atingimento de cota",
    }
    for i, (view, sql) in enumerate(kpi_sql().items(), 1):
        a.quadro_sql(f"SQL do KPI {i:02d} – {nomes[view]}", sql)

    yoy = n["yoy"]
    a.tabela("Resultado consolidado dos indicadores",
             ["Nº", "Indicador", "Resultado"],
             [
                 ["1", "Receita líquida (período total)", moeda(r.receita_liquida)],
                 ["2", "Pedidos / clientes ativos", f"{br(r.qtd_pedidos, 0)} pedidos / {br(r.clientes_ativos, 0)} clientes"],
                 ["3", "Ticket médio", f"US$ {br(r.ticket_medio)} (online 2025: US$ {br(tk_on[2025])}; revenda 2025: US$ {br(tk_rv[2025])})"],
                 ["4", "Margem bruta", f"{br(r.margem_bruta_pct)}% (online {br(mc.loc['Online','m'])}%; revenda {br(mc.loc['Revenda','m'])}%)"],
                 ["5", "Crescimento anual 2024 × 2023", f"{br(cresc_24, 1)}% (meses com crescimento em 2024: {int((yoy[yoy.ano == 2024].crescimento_yoy_pct > 0).sum())} de 12)"],
                 ["6", "Maior território", f"{terr.iloc[0].nome_territorio}: {br(terr.iloc[0].participacao_pct)}% da receita"],
                 ["7", "Maior categoria", f"Bikes: {br(cat.loc['Bikes','participacao_pct'])}% da receita, margem {br(cat.loc['Bikes','margem_bruta_pct'])}%"],
                 ["8", "Desconto médio", f"{br(r.desconto_medio_pct)}% do valor bruto"],
                 ["9", "Taxa de recompra", f"Lojas {br(rec.loc['Loja','taxa_recompra_pct'])}%; pessoas físicas {br(rec.loc['Pessoa física','taxa_recompra_pct'])}%"],
                 ["10", "Atingimento médio de cota", f"2023: {br(cota.loc[2023,'media'],1)}%; 2024: {br(cota.loc[2024,'media'],1)}%"],
             ], [1.0, 5.0, 9.5])

    a.titulo("2.7 ANÁLISE DOS RESULTADOS E STORYTELLING COM DADOS", 2)
    a.par(
        "Knaflic (2018) recomenda que a apresentação de dados comece pelo contexto, destaque o que é "
        "relevante e conduza o público a uma conclusão. Seguindo essa orientação, o dashboard foi organizado "
        "como uma narrativa em três atos: crescimento, rentabilidade e pontos de atenção. Os gráficos a seguir "
        "foram gerados a partir das mesmas views que alimentam o dashboard, cujo roteiro de montagem está no "
        "repositório (docs/dashboard.md).")
    a.titulo("2.7.1 Ato 1: uma empresa em crescimento, puxada pelo canal online", 3)
    a.par(
        f"A receita anual passou de {moeda(anual[2023])} em 2023 para {moeda(anual[2024])} em 2024, "
        f"crescimento de {br(cresc_24, 1)}%. Os anos de 2022 e 2025 aparecem parciais na base e não devem "
        f"ser comparados diretamente. A Figura 2 mostra que a revenda, de maior valor absoluto, oscila muito "
        f"de um mês para o outro, enquanto o canal online cresce de forma contínua a partir de meados de 2024. "
        f"A participação do online na receita subiu de {br(share_online(2022), 1)}% em 2022 para "
        f"{br(share_online(2025), 1)}% em 2025. Os últimos pedidos de revenda da base são de "
        f"{n['fim_revenda']:%m/%Y}, o que explica a queda desse canal no fim da série.")
    a.figura("fig2_receita_mensal_canal.png", "Receita líquida mensal por canal de venda")
    a.par(
        f"O crescimento online, porém, veio com mudança de perfil. O ticket médio online caiu de "
        f"US$ {br(tk_on[2022], 0)} em 2022 para US$ {br(tk_on[2025], 0)} em 2025 (Figura 3), porque a loja "
        "virtual passou a vender acessórios e roupas, itens baratos que aumentam o número de pedidos mas "
        "reduzem o valor médio de cada um. Esse é um padrão clássico de diversificação de mix.")
    a.figura("fig6_ticket_medio_online.png", "Ticket médio do canal online por ano")

    a.titulo("2.7.2 Ato 2: onde a empresa realmente ganha dinheiro", 3)
    a.par(
        f"O principal achado do projeto está na margem bruta (Figura 4). O canal online opera com margem "
        f"estável próxima de {br(mc.loc['Online','m'], 0)}%, enquanto a revenda caiu de {br(m_rv[2022], 1)}% em "
        f"2022 para {br(m_rv[2024], 1)}% em 2024, ou seja, passou a vender abaixo do custo padrão. Embora a "
        f"revenda responda por {br(100 * mc.loc['Revenda','r'] / r.receita_liquida, 0)}% da receita, é o canal "
        "online que gera a maior parte do lucro bruto. A revisão da política de preços para revendedores é, "
        "portanto, a ação de maior impacto sugerida pelos dados.")
    a.figura("fig3_margem_canal.png", "Margem bruta por canal de venda e ano")
    a.par(
        f"A concentração do portfólio reforça esse diagnóstico. Bicicletas representam "
        f"{br(cat.loc['Bikes','participacao_pct'], 1)}% da receita, com margem de "
        f"{br(cat.loc['Bikes','margem_bruta_pct'], 1)}%, enquanto acessórios, com apenas "
        f"{br(cat.loc['Accessories','participacao_pct'], 1)}% da receita, têm margem de "
        f"{br(cat.loc['Accessories','margem_bruta_pct'], 1)}% (Figura 5). Estimular a venda de acessórios junto "
        "com bicicletas é uma forma de elevar a rentabilidade média.")
    a.figura("fig4_receita_categoria.png", "Receita líquida e participação por categoria de produto")
    a.par(
        f"Geograficamente, a América do Norte concentra {br(grupos['North America'], 1)}% da receita, e os três "
        f"maiores territórios (Southwest, Canada e Northwest) somam {br(top3, 1)}% (Figura 6). Europa e "
        f"Pacífico, com {br(grupos['Europe'], 1)}% e {br(grupos['Pacific'], 1)}%, representam oportunidade de expansão.")
    a.figura("fig5_receita_territorio.png", "Receita líquida e participação por território de venda")

    a.titulo("2.7.3 Ato 3: clientes, promoções e equipe comercial", 3)
    a.par(
        f"A taxa de recompra revela dois comportamentos distintos: {br(rec.loc['Loja','taxa_recompra_pct'], 1)}% "
        f"das lojas compraram mais de uma vez, com média de {br(rec.loc['Loja','media_pedidos_por_cliente'], 1)} "
        f"pedidos cada, contra {br(rec.loc['Pessoa física','taxa_recompra_pct'], 1)}% das pessoas físicas. "
        "Programas de fidelização no canal online podem aumentar essa taxa. Os descontos têm impacto "
        f"pequeno, pois o desconto médio foi de {br(r.desconto_medio_pct)}% do valor bruto e se concentrou na "
        "revenda. Assim, a baixa margem desse canal decorre do preço de tabela para revendedores, e não das promoções.")
    a.par(
        f"Por fim, o atingimento médio de cota dos vendedores caiu de {br(cota.loc[2023,'media'], 1)}% em 2023 "
        f"para {br(cota.loc[2024,'media'], 1)}% em 2024. Em 2023, {int(cota.loc[2023,'bateram'])} de "
        f"{int(cota.loc[2023,'total'])} metas trimestrais foram batidas; em 2024, apenas "
        f"{int(cota.loc[2024,'bateram'])} de {int(cota.loc[2024,'total'])}. Somado à queda de margem, o dado "
        "indica que a equipe de revenda perdeu eficiência e deve ser objeto de revisão de metas e incentivos.")

    # ---------------- 3 Considerações finais
    a.titulo("3 CONSIDERAÇÕES FINAIS")
    a.par(
        "Este trabalho construiu um Data Warehouse de vendas para a AdventureWorks, desde a avaliação do "
        "modelo transacional até a entrega de indicadores prontos para consumo em um dashboard. O modelo em "
        "esquema estrela, com duas tabelas fato e sete dimensões conformadas, substituiu consultas que na origem "
        "exigiam muitas junções entre esquemas por consultas simples sobre uma fato e poucas dimensões, legíveis por "
        "analistas de negócio.")
    a.par(
        "A ETL em Python mostrou-se adequada ao propósito: é modular, configurável por variáveis de ambiente, "
        "reprodutível, validada contra a origem e executada em poucos segundos. Regras como o custo histórico "
        "e o rateio de frete e impostos agregam valor que não existia de forma direta no sistema transacional.")
    a.par(
        "Os dez indicadores contaram uma história clara: a empresa cresce, puxada pelo canal online, mas a "
        "revenda, responsável pela maior parte da receita, praticamente não gera margem e tem atingimento de "
        "metas em queda. Esse tipo de conclusão, difícil de obter diretamente do banco transacional, "
        "evidencia o papel do DW no apoio à decisão.")
    a.par(
        "Como trabalhos futuros, sugere-se implementar carga incremental com dimensões de variação lenta do "
        "tipo 2 (por exemplo, para preservar o histórico de território dos vendedores), incluir novos "
        "processos de negócio, como compras e produção, compartilhando as dimensões já criadas, e orquestrar "
        "a ETL com uma ferramenta de agendamento, como o Apache Airflow.")

    # ---------------- Referências
    a.titulo("REFERÊNCIAS")
    refs = [
        "CENTRO UNIVERSITÁRIO SALESIANO. **Guia de elaboração e normalização de trabalhos acadêmicos e de pesquisa**. Vitória: Unisales, 2024. Disponível em: https://unisales.br/wp-content/uploads/2024/07/NOVO-GUIA-DE-ELABORACAO-E-NORMALIZACAO-DE-TRABALHOS-ACADEMICOS-E-DE-PESQUISA-29.05.pdf. Acesso em: " + DATA_ACESSO + ".",
        "INMON, William H. **Building the data warehouse**. 4. ed. Indianapolis: Wiley, 2005.",
        "KIMBALL, Ralph; CASERTA, Joe. **The data warehouse ETL toolkit**: practical techniques for extracting, cleaning, conforming, and delivering data. Indianapolis: Wiley, 2004.",
        "KIMBALL, Ralph; ROSS, Margy. **The data warehouse toolkit**: the definitive guide to dimensional modeling. 3. ed. Indianapolis: Wiley, 2013.",
        "KNAFLIC, Cole Nussbaumer. **Storytelling com dados**: um guia sobre visualização de dados para profissionais de negócios. Rio de Janeiro: Alta Books, 2018.",
        "MICROSOFT. **Bancos de dados de exemplo AdventureWorks**. 2024. Disponível em: https://learn.microsoft.com/pt-br/sql/samples/adventureworks-install-configure?view=sql-server-ver17. Acesso em: " + DATA_ACESSO + ".",
        "POSTGRESQL GLOBAL DEVELOPMENT GROUP. **PostgreSQL 16 documentation**. 2024. Disponível em: https://www.postgresql.org/docs/16/. Acesso em: " + DATA_ACESSO + ".",
        "THWAITS, Lorin. **AdventureWorks for Postgres**. 2024. Disponível em: https://github.com/lorint/AdventureWorks-for-Postgres. Acesso em: " + DATA_ACESSO + ".",
    ]
    for ref in refs:
        a.par(ref, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=0)
        a.vazio()

    # Notas de autoria (rodapé simulado ao fim, Arial 10)
    a.par("¹ Graduando(a) em [curso] pelo Centro Universitário Salesiano (Unisales). E-mail: [e-mail].", tamanho=10, depois=0)
    a.par("² Professor(a) orientador(a) do Centro Universitário Salesiano (Unisales). E-mail: [e-mail].", tamanho=10, depois=0)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--github", default="https://github.com/SEU-USUARIO/adventureworks-dw")
    args = ap.parse_args()
    numeros = carregar_numeros(dw_engine())
    artigo = Artigo()
    escrever(artigo, numeros, args.github)
    artigo.salvar()


if __name__ == "__main__":
    main()
